from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from geom import *
from guide import GUIDE
from quantities import take_off

F = "Arial"
HEAD = PatternFill("solid", fgColor="2F6B4F")
INPUT = PatternFill("solid", fgColor="FFFF00")
TOTAL = PatternFill("solid", fgColor="E8EFEA")
thin = Side(style="thin", color="C8C8C8")
BOX = Border(bottom=thin)


def sheet(wb, title, headers, widths):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
        c = ws.cell(1, i)
        c.font = Font(name=F, bold=True, color="FFFFFF"); c.fill = HEAD
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"
    return ws


def style_rows(ws, ncols):
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=ncols):
        for c in row:
            if c.font is None or c.font.name != F or not c.font.bold:
                c.font = Font(name=F, size=10, bold=c.font.bold if c.font else False)
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = BOX


def kind(item):
    i = item.lower()
    if any(k in i for k in ("wc", "basin", "shower", "sink")): return "Sanitary / plumbing"
    if any(k in i for k in ("refrigerator", "washing machine", "tv unit")): return "Appliance / joinery"
    if any(k in i for k in ("wardrobe", "counter", "shelving", "bookshelf", "cupboard", "arugu", "crockery", "mandapam", "airer")): return "Built-in joinery"
    return "Loose furniture"


def make(path):
    wb = Workbook()
    rd = wb.active; rd.title = "Read_Me"
    lines = [
        ("Kerala Home, Phase 1: schedules", True),
        (f"Site: {PROJECT['location']}. Stage: {PROJECT['stage']}. Revision {PROJECT['rev']}, {PROJECT['date']}.", False),
        ("Dimensions in feet. Centreline sizes match the plan grid; clear sizes are wall face to wall face after plaster.", False),
        ("", False),
        ("How to use", True),
        ("Yellow cells are for you to fill in (budget estimates, chosen products, status). Everything else is calculated or drawn from the plan.", False),
        ("Example: in Furniture, type 18000 in 'Est. unit cost (Rs)' for an item costing Rs 18,000; the line total and grand total update automatically.", False),
        ("Areas, heights and totals are formulas, so editing a room's size updates its area.", False),
        ("", False),
        ("Tabs", True),
        ("Rooms: room sizes and areas.  Doors / Windows: opening sizes and types.  Furniture: every furniture item, fixture and built-in, with budget columns.", False),
        ("Finishes: floor, wall and ceiling finishes.  Specification: materials and why (Guntur conditions).  Quantities: take-off for pricing, with rate cells.", False),
        ("Electrical / Plumbing: every point with room and height.  Checklist: room-by-room mistakes to avoid, with a status column.  Hold_Points: inspections before work is covered.", False),
        ("Open_Issues: items to resolve with professionals.  Decisions: what has been agreed so far.", False),
    ]
    for t, b in lines:
        rd.append([t]); rd.cell(rd.max_row, 1).font = Font(name=F, bold=b, size=13 if b else 10)
    rd.column_dimensions["A"].width = 120
    for r in rd.iter_rows():
        r[0].alignment = Alignment(wrap_text=True, vertical="top")

    ws = sheet(wb, "Rooms", ["ID", "Room", "Zone", "E-W (ft)", "N-S (ft)", "Area (sq ft)", "Area (sq m)", "Clear E-W (ft)", "Clear N-S (ft)", "Clear area (sq ft)", "Floor level", "Ceiling", "Notes"],
               [7, 24, 14, 9, 9, 11, 11, 10, 10, 11, 11, 16, 34])
    notes = {"R10": "Under the ventilator lantern, at plot centre", "R13": "Semi-open, jali on south wall", "R17": "Open verandah", "R09": "Includes 2' x 6' built-in wardrobe",
             "R04": "Widened 1'-6\" into the master: WC, 3'-6\" vanity, walk-in shower"}
    for i, r in enumerate(ROOMS, start=2):
        L, W = room_dims(r)
        cl, cw = clear_dims(r)
        ck, cv = CEILINGS[r[0]]
        ws.append([r[0], r[1], r[6], L, W, f"=D{i}*E{i}", f"=F{i}*0.092903", round(cl, 2), round(cw, 2), f"=H{i}*I{i}",
                   fmt_lvl(floor_level(r[0])), f"Flat {fmt_ft(cv)}" if ck == "flat" else f"Sloped ({cv})", notes.get(r[0], "")])
    n = ws.max_row
    ws.append(["", "Total", "", "", "", f"=SUM(F2:F{n})", f"=SUM(G2:G{n})", "", "", f"=SUM(J2:J{n})", "", "", "Built footprint 60' x 22' = 1,320 sq ft"])
    style_rows(ws, 13)
    for c in ws[ws.max_row]: c.font = Font(name=F, bold=True); c.fill = TOTAL
    for row in ws.iter_rows(min_row=2, min_col=6, max_col=10):
        for c in row: c.number_format = "#,##0.0"
    ws.cell(1, 7).comment = Comment("1 sq ft = 0.092903 sq m", "Claude")

    ws = sheet(wb, "Doors", ["ID", "Location", "Width (ft)", "Height (ft)", "Operation", "Type / material", "Vent above (7'-6\" to 9'-0\")", "Chosen product", "Status"], [7, 26, 10, 10, 14, 40, 44, 26, 14])
    for d in DOORS:
        op = "Sliding" if d[5] == 0 else "Hinged, single" if "double" not in d[7] else "Hinged, double"
        ws.append([d[0], d[6], round(d[4] - d[3], 2), 7, op, d[7], TRANSOM_TYPES[TRANSOMS[d[0]]] if d[0] in TRANSOMS else "-", "", ""])
    style_rows(ws, 9)
    for row in ws.iter_rows(min_row=2, min_col=8, max_col=9):
        for c in row: c.fill = INPUT

    ws = sheet(wb, "Windows", ["ID", "Room", "Wall faces", "Width (ft)", "Sill (ft)", "Head (ft)", "Height (ft)", "Type", "Status"], [7, 18, 10, 10, 9, 9, 10, 52, 14])
    face = {("X", 3): "South", ("X", 25): "North", ("Y", 5): "West", ("Y", 65): "East", ("Y", 59): "East (to verandah)"}
    for i, w in enumerate(WINDOWS, start=2):
        ws.append([w[0], w[7], face.get((w[1], w[2]), ""), round(w[4] - w[3], 2), w[5], w[6], f"=F{i}-E{i}", w[8], ""])
    for o in OTHER_OPENINGS:
        ws.append([o[0], "Roof", "Up", "", "", "", "", o[1], ""])
    style_rows(ws, 9)
    for row in ws.iter_rows(min_row=2, min_col=9, max_col=9):
        for c in row: c.fill = INPUT
    for row in ws.iter_rows(min_row=2, min_col=4, max_col=7):
        for c in row: c.number_format = "0.0#"

    ws = sheet(wb, "Furniture", ["ID", "Room", "Item", "Category", "E-W (ft)", "N-S (ft)", "Qty", "Material / spec", "Est. unit cost (Rs)", "Line total (Rs)", "Chosen product / vendor", "Status"],
               [7, 18, 34, 18, 9, 9, 6, 34, 14, 14, 26, 12])
    for i, f in enumerate(FURNITURE, start=2):
        ws.append([f[0], f[1], f[2], kind(f[2]), round(f[4] - f[3], 2), round(f[6] - f[5], 2), f[7], f[8], None, f"=G{i}*I{i}", "", ""])
    n = ws.max_row
    ws.append(["", "", "Total", "", "", "", f"=SUM(G2:G{n})", "", "", f"=SUM(J2:J{n})", "", ""])
    style_rows(ws, 12)
    for c in ws[ws.max_row]: c.font = Font(name=F, bold=True); c.fill = TOTAL
    for row in ws.iter_rows(min_row=2, max_row=n, min_col=9, max_col=12):
        for c in row:
            if c.column in (9, 11, 12): c.fill = INPUT
    for row in ws.iter_rows(min_row=2, min_col=9, max_col=10):
        for c in row: c.number_format = "#,##0;(#,##0);-"
    ws.cell(1, 9).comment = Comment("Your estimate per unit, in rupees. Leave blank until you have a quote.", "Claude")

    ws = sheet(wb, "Finishes", ["Space", "Floor", "Walls", "Ceiling", "Notes", "Decision"], [34, 32, 32, 34, 26, 16])
    for f in FINISHES: ws.append(list(f) + [""])
    style_rows(ws, 6)
    for row in ws.iter_rows(min_row=2, min_col=6, max_col=6):
        for c in row: c.fill = INPUT

    ws = sheet(wb, "Specification", ["Category", "Element", "Specification", "Why (Guntur)"], [14, 24, 70, 50])
    for row in SPEC: ws.append(list(row))
    style_rows(ws, 4)

    ws = sheet(wb, "Quantities", ["Group", "Item", "Qty", "Unit", "Basis", "Rate (Rs / unit)", "Amount (Rs)"], [12, 46, 10, 8, 46, 14, 16])
    for i, (g, item, qty, unit, basis) in enumerate(take_off(), start=2):
        ws.append([g, item, qty, unit, basis, None, f"=IF(F{i}=\"\",\"\",C{i}*F{i})"])
    n = ws.max_row
    ws.append(["", "Total of priced items", "", "", "", "", f"=SUM(G2:G{n})"])
    style_rows(ws, 7)
    for c in ws[ws.max_row]: c.font = Font(name=F, bold=True); c.fill = TOTAL
    for row in ws.iter_rows(min_row=2, max_row=n, min_col=6, max_col=6):
        for c in row: c.fill = INPUT
    for row in ws.iter_rows(min_row=2, min_col=6, max_col=7):
        for c in row: c.number_format = "#,##0;(#,##0);-"
    ws.cell(1, 6).comment = Comment("Fill from local contractor or supplier quotes. Quantities are approximate take-offs from the model.", "Claude")

    ws = sheet(wb, "Electrical", ["Code", "Item", "Mounting height", "Room", "X (ft)", "Y (ft)", "Note"], [7, 34, 14, 22, 8, 8, 40])
    rooms = {r[0]: r[1] for r in ROOMS} | {"SITE": "Site"}
    for code, x, y, rid, note in ELECTRICAL:
        ws.append([code, ELEC_TYPES[code][0], ELEC_TYPES[code][1], rooms[rid], x, y, note])
    style_rows(ws, 7)

    ws = sheet(wb, "Plumbing", ["Code", "Item", "Room", "X (ft)", "Y (ft)", "Note"], [7, 34, 22, 8, 8, 40])
    for code, x, y, rid, note in PLUMBING:
        ws.append([code, PLUMB_TYPES[code], rooms[rid], x, y, note])
    style_rows(ws, 6)

    ws = sheet(wb, "Checklist", ["Area", "Common mistake", "Do instead", "Shown on", "Status"], [22, 50, 60, 14, 14])
    for area, items in GUIDE:
        for m, fix, ref in items:
            ws.append([area, m, fix, ref, ""])
    style_rows(ws, 5)
    for row in ws.iter_rows(min_row=2, min_col=5, max_col=5):
        for c in row: c.fill = INPUT
    ws.cell(1, 5).comment = Comment("Mark Agreed / Changed / Not needed after reviewing with your architect.", "Claude")

    ws = sheet(wb, "Hold_Points", ["#", "Stage", "Check before it is covered up", "Who signs off", "Date checked", "Photos / notes"], [5, 22, 70, 24, 14, 30])
    for i, h in enumerate(HOLD_POINTS, 1): ws.append([i] + list(h) + ["", ""])
    style_rows(ws, 6)
    for row in ws.iter_rows(min_row=2, min_col=5, max_col=6):
        for c in row: c.fill = INPUT

    ws = sheet(wb, "Open_Issues", ["#", "Category", "Issue", "Who resolves", "Status"], [5, 12, 80, 22, 18])
    for i, o in enumerate(OPEN_ISSUES, 1): ws.append([i] + list(o))
    style_rows(ws, 5)
    for row in ws.iter_rows(min_row=2, min_col=5, max_col=5):
        for c in row: c.fill = INPUT

    ws = sheet(wb, "Decisions", ["#", "Decision (agreed in concept stage)"], [5, 110])
    for i, d in enumerate(DECISIONS, 1): ws.append([i, d])
    style_rows(ws, 2)
    wb.save(path)


if __name__ == "__main__":
    from paths import XLSX
    make(str(XLSX))
